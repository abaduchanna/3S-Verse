#!/usr/bin/env python3
# ============================================================
#  GFH INVENTORY DASHBOARD - FIREBASE UPLOADER v3.8 (MULTI-CREDENTIAL AUTH + ZERO-PROMPT)
# ============================================================
#  Kya karta hai:
#    Aap ki Excel (.xlsx) ya CSV file ke SAARE rows read kar ke
#    dashboard (gfhinventorydashboard.netlify.app) ke Firebase
#    database par upload kar deta hai.
#
#  CREDENTIAL ENGINE (v3.8 - file ka TYPE khud pehchanta hai):
#    Chaar types support hain - jo bhi credential.json / credentials.json
#    mile, us ke andar ki keys se type detect hota hai:
#      1. Firebase SERVICE-ACCOUNT key (private_key + client_email)
#         -> RSA JWT se Google access token (pure Python - koi extra
#            install nahi, sirf openpyxl pehle jaisa)
#      2. Google ADC / gcloud credential (refresh_token + client_id)
#         -> refresh flow se access token
#      2b. Google OAuth CLIENT file ({"installed": {...}} / gcloud
#          client-secret - client_id + client_secret, refresh_token
#          NAHI) -> EK DAFA browser consent flow (aap Google login
#          karo, Allow dabao) -> credential_user.json me save ->
#          agli baar browser nahi khulega
#      3. Firebase WEB CONFIG (apiKey) -> anonymous sign-in
#         (Console > Authentication > Anonymous ON hona chahiye)
#    Search: is folder + Desktop + Downloads + Documents + 2-level
#    subfolders. credential.json / credentials.json / .txt variants /
#    Notepad (BOM/UTF-16/ANSI) sab chalte hain. Milti hai to copy is
#    folder mein bhi bana deta hai (agle run instant).
#
#  FILE PICK (bina scan circus):
#    - Pichli baar jo file upload hui thi, wohi agle run par seedha
#      uthata hai (LAST_UPLOAD.txt memory - koi 23-file scan NAHI).
#    - Memory na ho to seedha 'gfh database.xlsx' (.xlsm/.csv bhi).
#    - Wo bhi na miley to chup-chaap best-match scan (sirf winner
#      print hota hai, poori list nahi).
#    - Drag & drop wala tareeqa pehle jaisa hai.
#
#  Zaroori:
#    - 'gfh database.xlsx' (tab: database, 25 columns) seedha chalti hai.
#    - KOI SAWAL / ENTER / yes-no NAHI: file mili to foran purana data
#      clear + naya upload. Bas RUN_UPLOAD.bat chalao.
#    - Columns ka ORDER ab farak nahi parta: script columns ko NAAM
#      se pehchan kar dashboard ke order mein khud set karti hai.
#    - Upload pehle PURANA DATA CLEAR (DELETE), phir POORA data EK hi
#      PUT mein charhta hai - dashboard isi shape ko parhta hai
#      (header + rows wala array). Pehli baar RUN_BACKUP.bat chalao.
#
#  Use:
#    RUN_UPLOAD.bat double-click karo
#    Ya Excel file ko is .bat par drag & drop kar do.
# ============================================================

import json
import os
import sys
import csv
import time
import datetime

# ------------------------- CONFIG ---------------------------
FIREBASE_DB_URL = "https://gfh-inventory-dashboard-6febd-default-rtdb.firebaseio.com"
PROJECT_ID = "gfh-inventory-dashboard-6febd"
DATA_NODE = "database"          # dashboard isi node se data uthata hai
SHEET_NAME = "database"         # data isi tab se uthaya jayega; na miley to headers se khud dhoondega
AUTO_PICK_MIN = 15              # auto-pick: kam az kam itne/25 headers match hon (Rebate-type report ~13 hai, wo refuse ho jati hai)
# ------------------------------------------------------------

EXPECTED_HEADERS = [
    "District", "Store name", "Date", "Order details", "Product description",
    "ESN number", "Tracking number", "Tracking status", "Cost", "Due date",
    "Payment due status", "Stock status", "Ext Price", "Expected Commission",
    "ER Exp Comm", "Rebate", "SP Exp Comm", "1st Month Spiff",
    "Rebate variance", "Spiff Variance", "Device missing status",
    "Device sale location", "Device sale date", "Device sold by",
    "External Order ID",
]

LINE = "-" * 58


def ensure_package(modname, pipname):
    try:
        __import__(modname)
        return True
    except ImportError:
        print(f"[SETUP] {modname} missing hai - install kar raha hoon (internet chahiye, ~20 sec)...")
        try:
            import subprocess
            subprocess.check_call([sys.executable, "-m", "pip", "install", pipname])
            __import__(modname)
            return True
        except Exception as e:
            print(f"[ERROR] {modname} install nahi ho saka: {e}")
            print(f"        Manually: pip install {pipname}")
            return False


def ensure_openpyxl():
    return ensure_package("openpyxl", "openpyxl")


# ---------------- AUTH ENGINE (v3.7) ----------------
# Credential types jo file content se khud detect hote hain:
#   service_account : private_key + client_email (Firebase Admin key)
#   authorized_user : refresh_token + client_id + client_secret (gcloud ADC)
#   web_config      : apiKey (Firebase web app config) -> anonymous sign-in
# Tokens PURE Python bante hain (RSA PKCS#1 v1.5 + SHA-256) - koi
# extra pip install nahi chahiye.

CRED_LABELS = {
    "service_account": "Firebase SERVICE-ACCOUNT key detect hui",
    "authorized_user": "Google ADC / gcloud credential detect hui",
    "desktop_client": "Google OAuth CLIENT file detect hui (browser login SIRF EK DAFA lagega)",
    "web_config": "Firebase WEB CONFIG detect hui",
}


def classify_cred(d):
    """dict -> credential type key ya None."""
    if not isinstance(d, dict):
        return None
    if (str(d.get("private_key") or "").strip()
            and str(d.get("client_email") or "").strip()):
        return "service_account"
    if (str(d.get("refresh_token") or "").strip()
            and str(d.get("client_id") or "").strip()
            and str(d.get("client_secret") or "").strip()):
        return "authorized_user"
    # gcloud / Google Cloud Console wali OAuth client-secret file:
    # {"installed": {client_id, client_secret, auth_uri, ...}}
    if (str(d.get("client_id") or "").strip()
            and str(d.get("client_secret") or "").strip()
            and (str(d.get("auth_uri") or "").strip()
                 or str(d.get("token_uri") or "").strip()
                 or str(d.get("redirect_uris") or "").strip())):
        return "desktop_client"
    if str(d.get("apiKey") or "").strip():
        return "web_config"
    return None


def unwrap_cred(d):
    """gcloud client-secret files {'installed':{...}} / {'web':{...}} kholta hai."""
    if isinstance(d, dict):
        if isinstance(d.get("installed"), dict):
            return d["installed"]
        if isinstance(d.get("web"), dict):
            return d["web"]
    return d


def _der_read(buf, off):
    """Chhota ASN.1 DER TLV reader -> (tag, content, next_offset)."""
    tag = buf[off]
    off += 1
    ln = buf[off]
    off += 1
    if ln & 0x80:
        nb = ln & 0x7F
        ln = int.from_bytes(buf[off:off + nb], "big")
        off += nb
    return tag, buf[off:off + ln], off + ln


def _b64u(b):
    import base64
    return base64.urlsafe_b64encode(b).rstrip(b"=").decode("ascii")


def _rsa_from_pem(pem_text):
    """'BEGIN PRIVATE KEY' (PKCS#8) ya 'BEGIN RSA PRIVATE KEY' (PKCS#1)
    PEM se (n, e, d) - pure Python."""
    import base64, re

    def _rsa_ints(fields):
        """RSAPrivateKey SEQUENCE ka content -> (n, e, d)."""
        ints, off2 = [], 0
        for _ in range(4):            # version, n, e, d
            tag, v, off2 = _der_read(fields, off2)
            if tag != 0x02:
                raise ValueError("RSAPrivateKey INTEGER nahi mila")
            ints.append(int.from_bytes(v, "big"))
        return ints[1], ints[2], ints[3]

    m = re.search(r"-----BEGIN (?:RSA )?PRIVATE KEY-----(.*?)-----END (?:RSA )?PRIVATE KEY-----",
                  pem_text, re.S)
    if not m:
        raise ValueError("PEM private key block nahi mila")
    der = base64.b64decode("".join(m.group(1).split()))
    tag, body, _ = _der_read(der, 0)
    if tag != 0x30:
        raise ValueError("DER SEQUENCE nahi mila")
    tag, ver, off = _der_read(body, 0)
    if tag == 0x02 and ver == b"\x00" and off < len(body):
        tag, nxt, off = _der_read(body, off)
        if tag == 0x30:
            # PKCS#8: version, alg SEQUENCE, OCTET STRING(RSAPrivateKey)
            tag, inner, off = _der_read(body, off)
            if tag == 0x04:
                tag, seq, _ = _der_read(inner, 0)
                if tag != 0x30:
                    raise ValueError("RSAPrivateKey SEQUENCE nahi mili")
                return _rsa_ints(seq)
    # PKCS#1: body khud RSAPrivateKey fields ka content hai
    return _rsa_ints(body)


def _rs256_sign(signing_input, n, e, d):
    """RS256 (PKCS#1 v1.5 + SHA-256) - pure Python."""
    import hashlib
    digest = hashlib.sha256(signing_input).digest()
    prefix = bytes.fromhex("3031300d060960864801650304020105000420")
    di = prefix + digest
    k = (n.bit_length() + 7) // 8
    padded = b"\x00\x01" + b"\xff" * (k - len(di) - 3) + b"\x00" + di
    if len(padded) != k:
        raise ValueError("RSA padding size galat")
    sig = pow(int.from_bytes(padded, "big"), d, n)
    return sig.to_bytes(k, "big")


def _http_post(url, data_bytes, content_type, timeout=60):
    """POST -> (status, parsed_json). HTTPError body bhi parse hota hai."""
    import urllib.request
    import urllib.error
    req = urllib.request.Request(url, data=data_bytes, headers={
        "Content-Type": content_type})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return resp.status, json.loads(resp.read() or b"{}")
    except urllib.error.HTTPError as e:
        try:
            return e.code, json.loads(e.read() or b"{}")
        except Exception:
            return e.code, {}
    except Exception as e:
        raise RuntimeError(f"Network fail ({url}): {str(e)[:150]}")


def sa_mint_token(cred):
    """Service-account JSON -> OAuth2 access_token (RTDB scope). RSA JWT."""
    import urllib.parse
    n, e, d = _rsa_from_pem(str(cred["private_key"]))
    now = int(time.time())
    hdr = _b64u(json.dumps({"alg": "RS256", "typ": "JWT"},
                           separators=(",", ":")).encode())
    claims = {
        "iss": cred["client_email"],
        "scope": ("https://www.googleapis.com/auth/firebase.database "
                  "https://www.googleapis.com/auth/userinfo.email"),
        "aud": "https://oauth2.googleapis.com/token",
        "iat": now,
        "exp": now + 3600,
    }
    pl = _b64u(json.dumps(claims, separators=(",", ":")).encode())
    sig = _rs256_sign(f"{hdr}.{pl}".encode(), n, e, d)
    st, resp = _http_post(
        "https://oauth2.googleapis.com/token",
        urllib.parse.urlencode({
            "grant_type": "urn:ietf:params:oauth:grant-type:jwt-bearer",
            "assertion": f"{hdr}.{pl}.{_b64u(sig)}",
        }).encode(),
        "application/x-www-form-urlencoded")
    if st != 200 or "access_token" not in resp:
        raise RuntimeError(f"Google token exchange fail (HTTP {st}): {str(resp)[:200]}")
    return resp["access_token"]


def adc_mint_token(cred):
    """gcloud ADC (authorized_user) -> access_token via refresh_token."""
    import urllib.parse
    st, resp = _http_post(
        str(cred.get("token_uri") or "https://oauth2.googleapis.com/token"),
        urllib.parse.urlencode({
            "grant_type": "refresh_token",
            "refresh_token": cred["refresh_token"],
            "client_id": cred["client_id"],
            "client_secret": cred["client_secret"],
        }).encode(),
        "application/x-www-form-urlencoded")
    if st != 200 or "access_token" not in resp:
        raise RuntimeError(f"Google refresh fail (HTTP {st}): {str(resp)[:200]}")
    return resp["access_token"]


def desktop_login_flow(client, script_dir, log=print):
    """Google OAuth CLIENT file (gcloud client-secret / 'installed' shape)
    -> EK DAFA browser consent -> token exchange -> credential_user.json
    me save (authorized_user shape). Agle runs ke liye browser NAHI
    khulega - seedha refresh_token se token banta hai.
    Returns: (fresh access_token, saved_json_path)"""
    import http.server
    import socket as _sock
    import threading
    import urllib.parse
    import webbrowser

    # free loopback port (desktop OAuth clients allow any 127.0.0.1 port)
    s = _sock.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    redirect_uri = f"http://127.0.0.1:{port}"
    scope = ("https://www.googleapis.com/auth/firebase.database "
             "https://www.googleapis.com/auth/userinfo.email")
    auth_base = (str(client.get("auth_uri") or "").strip()
                 or "https://accounts.google.com/o/oauth2/v2/auth")
    auth_url = auth_base + "?" + urllib.parse.urlencode({
        "client_id": client["client_id"],
        "redirect_uri": redirect_uri,
        "response_type": "code",
        "scope": scope,
        "access_type": "offline",
        "prompt": "consent",
    })

    holder = {"code": None, "error": None}

    class _H(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            q = urllib.parse.parse_qs(
                urllib.parse.urlparse(self.path).query)
            holder["code"] = (q.get("code") or [None])[0]
            holder["error"] = (q.get("error") or [None])[0]
            ok = holder["code"] is not None
            msg = ("Login OK - ye tab band kar do aur uploader par wapas "
                   "aao." if ok else
                   f"Login fail: {holder['error']} - uploader dobara chalao.")
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.end_headers()
            self.wfile.write(
                ("<html><body style='font-family:Segoe UI,sans-serif;"
                 "text-align:center;padding-top:70px'><h2>" + msg +
                 "</h2></body></html>").encode("utf-8"))

        def log_message(self, *a):
            pass

    srv = http.server.HTTPServer(("127.0.0.1", port), _H)
    threading.Thread(target=srv.handle_request, daemon=True).start()

    log("[AUTH] Browser khul raha hai - Google account chuno aur 'Allow' "
        "dabao (SIRF EK DAFA).")
    log("       Agar browser na khule to ye link khud kholo:")
    log(f"       {auth_url}")
    try:
        webbrowser.open(auth_url)
    except Exception:
        pass

    deadline = time.time() + 300.0
    while time.time() < deadline:
        if holder["code"] or holder["error"]:
            break
        time.sleep(0.5)
    try:
        srv.server_close()
    except Exception:
        pass
    if holder["error"]:
        raise RuntimeError(f"Google ne login mana kar diya: "
                           f"{holder['error']} - dobara chalao aur 'Allow' "
                           f"dabao.")
    if not holder["code"]:
        raise RuntimeError("Google login poora nahi hua (5 min timeout) - "
                           "dobara chalao aur 'Allow' dabao.")

    st, resp = _http_post(
        str(client.get("token_uri") or "https://oauth2.googleapis.com/token"),
        urllib.parse.urlencode({
            "code": holder["code"],
            "client_id": client["client_id"],
            "client_secret": client["client_secret"],
            "redirect_uri": redirect_uri,
            "grant_type": "authorization_code",
        }).encode(),
        "application/x-www-form-urlencoded")
    if st != 200 or "refresh_token" not in resp:
        raise RuntimeError(f"Google code exchange fail (HTTP {st}): "
                           f"{str(resp)[:200]}")

    adc = {
        "type": "authorized_user",
        "refresh_token": resp["refresh_token"],
        "client_id": client["client_id"],
        "client_secret": client["client_secret"],
        "token_uri": (str(client.get("token_uri") or "").strip()
                      or "https://oauth2.googleapis.com/token"),
    }
    dst = os.path.join(script_dir, "credential_user.json")
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(adc, f, indent=2)
    log("[OK] Login save ho gaya (credential_user.json) - agli baar "
        "browser NAHI khulega.")
    return resp["access_token"], dst


def anon_mint_token(webcfg):
    """Firebase WEB CONFIG -> anonymous idToken (Identity Toolkit).
    Chalta hai agar Anonymous provider ON hai aur rules 'auth != null'."""
    st, resp = _http_post(
        f"https://identitytoolkit.googleapis.com/v1/accounts:signUp"
        f"?key={webcfg['apiKey']}",
        json.dumps({"returnSecureToken": True}).encode(),
        "application/json")
    if st != 200 or "idToken" not in resp:
        err = (resp.get("error", {}).get("message", str(resp)[:120])
               if isinstance(resp, dict) else str(resp)[:120])
        raise RuntimeError(f"anonymous sign-in fail (HTTP {st}): {err}")
    return resp["idToken"]


def mint_token(cred, ctype):
    """Credential + type -> fresh auth token."""
    c = unwrap_cred(cred)
    if ctype == "service_account":
        return sa_mint_token(c)
    if ctype == "authorized_user":
        return adc_mint_token(c)
    if ctype == "web_config":
        return anon_mint_token(c)
    if ctype == "desktop_client":
        raise RuntimeError("desktop_client ko pehle desktop_login_flow se "
                           "login karna parta hai")
    raise RuntimeError(f"Credential type '{ctype}' support nahi hai")


class FBAuth:
    """RTDB REST auth: sa/adc -> ?access_token=..., anon -> ?auth=..."""

    def __init__(self, mode, token):
        self.mode, self.token = mode, token

    def url(self, base):
        import urllib.parse
        sep = "&" if "?" in base else "?"
        if self.mode in ("service_account", "authorized_user"):
            return f"{base}{sep}access_token={urllib.parse.quote(self.token, safe='')}"
        return f"{base}{sep}auth={urllib.parse.quote(self.token, safe='')}"


def cell_to_json(v):
    """Excel cell ko Firebase-friendly JSON value banata hai."""
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%dT%H:%M:%S.000")
    if isinstance(v, datetime.date):
        return v.strftime("%Y-%m-%dT00:00:00.000")
    if isinstance(v, datetime.time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, float):
        if v != v or v in (float("inf"), float("-inf")):   # NaN / inf
            return None
        return v
    if isinstance(v, str):
        return v.strip()
    return v


def is_empty_row(row):
    return all(v is None or (isinstance(v, str) and v.strip() == "") for v in row)


def to_number(s):
    """CSV text value ko number banata hai agar number jaisa lage."""
    if s is None:
        return None
    t = str(s).strip()
    if not t:
        return ""
    try:
        f = float(t)
        if f.is_integer() and "." not in t and "e" not in t.lower():
            return int(t)
        return f
    except ValueError:
        return t


def norm_first_row(ws, limit=300):
    """Sheet ki pehli row ke normalized (lower/strip) cell values."""
    try:
        for row in ws.iter_rows(min_row=1, max_row=1, max_col=limit, values_only=True):
            return [("" if v is None else str(v).strip().lower()) for v in row]
    except Exception:
        pass
    return []


def pick_sheet(wb):
    """Har sheet ke header row ko EXPECTED_HEADERS se score kar ke
    sab se matching sheet chunta hai. Returns (sheet_title, score)."""
    want_set = {h.lower() for h in EXPECTED_HEADERS}
    best, best_score = None, -1
    for ws in wb.worksheets:
        got = norm_first_row(ws)
        score = sum(1 for c in got if c in want_set)
        if score > best_score:
            best, best_score = ws.title, score
    return best, best_score


def read_xlsx(path):
    if not ensure_openpyxl():
        sys.exit(1)
    from openpyxl import load_workbook
    print(f"[READ] Excel khol raha hoon: {os.path.basename(path)}")
    wb = load_workbook(path, data_only=True, read_only=True)
    sheet_used = None
    if SHEET_NAME:
        for name in wb.sheetnames:
            if name.strip().lower() == SHEET_NAME.strip().lower():
                ws = wb[name]
                sheet_used = name
                break
        if sheet_used is None:
            print(f"[WARN] Sheet '{SHEET_NAME}' nahi mili (Sheets: {wb.sheetnames})")
            print("       Headers se match kar ke khud dhoond raha hoon...")
    if sheet_used is None:
        best, score = pick_sheet(wb)
        if best is None or score < 5:
            print("[ERROR] Koi sheet aisi nahi mili jis ke headers dashboard se match karein.")
            for ws in wb.worksheets:
                sc = sum(1 for c in norm_first_row(ws) if c in {h.lower() for h in EXPECTED_HEADERS})
                print(f"        - '{ws.title}': {sc}/25 headers match")
            wb.close()
            sys.exit(1)
        ws = wb[best]
        sheet_used = best
    print(f"[SHEET] '{sheet_used}' tab se data uthaya jayega.")
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append([cell_to_json(v) for v in r])
    wb.close()
    return rows


def read_csv(path):
    print(f"[READ] CSV khol raha hoon: {os.path.basename(path)}")
    rows = []
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        for raw in csv.reader(f):
            rows.append([cell_to_json(to_number(v)) for v in raw])
    return rows


def map_columns(header, data):
    """File columns ko dashboard ke 25-column order mein NAAM se map karta hai.
    - Order kuch bhi ho, data sahi jagah jata hai (naam match hona chahiye).
    - Jo dashboard columns file mein nahi wo khaali (None) jate hain.
    - File ke extra columns drop hote hain (report mein dikhte hain).
    Returns: (payload, matched_count, missing, extra)"""
    norm = lambda x: ("" if x is None else str(x).strip().lower())
    got = [norm(h) for h in header]
    want = [h.lower() for h in EXPECTED_HEADERS]
    want_set = set(want)

    index_by_name = {}
    for i, g in enumerate(got):
        if g and g not in index_by_name:
            index_by_name[g] = i

    matched = [w for w in want if w in index_by_name]
    missing = [w for w in want if w not in index_by_name]
    extra, seen = [], set()
    for g in got:
        if g and g not in want_set and g not in seen:
            extra.append(g)
            seen.add(g)

    ncols = len(header)
    mapped = []
    for r in data:
        r = list(r[:ncols]) + [None] * (ncols - len(r))
        mapped.append([r[index_by_name[w]] if w in index_by_name else None for w in want])

    payload = [list(EXPECTED_HEADERS)] + mapped
    return payload, len(matched), missing, extra


def build_payload(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        rows = read_csv(path)
    else:
        rows = read_xlsx(path)

    rows = [r for r in rows if not is_empty_row(r)]
    if len(rows) < 2:
        print("[ERROR] File mein header + kam az kam 1 data row hone chahiye.")
        sys.exit(1)

    header, data = rows[0], rows[1:]

    print("[MAP] Columns ko naam se dashboard ke order mein set kar raha hoon...")
    payload, nmatch, missing, extra = map_columns(header, data)

    if nmatch < 5:
        print(f"[ERROR] File ke headers dashboard se match nahi karte (sirf {nmatch}/25 mile).")
        print("        Ye shayad dashboard wali data file nahi hai.")
        fl = ", ".join(str(h) for h in header if h is not None)
        print(f"        File ke headers the: {fl[:300]}")
        sys.exit(1)

    print(f"[MAP] {nmatch}/25 columns file se map ho gaye.")
    if missing:
        ml = ", ".join(missing)
        print(f"      File mein nahi the ({len(missing)}) - khaali jayengi: {ml[:220]}{'...' if len(ml) > 220 else ''}")
    if extra:
        el = ", ".join(extra)
        print(f"      Extra drop hui ({len(extra)}): {el[:220]}{'...' if len(el) > 220 else ''}")
    return payload


def fb_request(url, method="GET", payload=None, timeout=600):
    """Firebase REST call. 4xx (401/403/400...) par retry NAHI - foran
    (status, resp) wapis. Network error / 5xx / 429 / 408 par 3 retries."""
    import urllib.request
    import urllib.error
    data = None
    headers = {}
    if payload is not None:
        data = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        headers["Content-Type"] = "application/json; charset=utf-8"
    last_err = None
    for attempt in range(1, 4):
        try:
            req = urllib.request.Request(url, data=data, headers=headers, method=method)
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                body = resp.read()
                try:
                    return resp.status, json.loads(body) if body else None
                except Exception:
                    return resp.status, body
        except urllib.error.HTTPError as e:
            body = b""
            try:
                body = e.read()
            except Exception:
                pass
            try:
                parsed = json.loads(body) if body else None
            except Exception:
                parsed = body
            if e.code in (408, 429) or 500 <= e.code < 600:
                last_err = e                     # retry-able
                if attempt < 3:
                    print(f"    HTTP {e.code} - retry {attempt}/3, 5 sec ruk kar...")
                    time.sleep(5)
                    continue
            return e.code, parsed                # 4xx: retry bekaar
        except Exception as e:
            last_err = e
            if attempt < 3:
                print(f"    Network error ({e}) - retry {attempt}/3, 5 sec ruk kar...")
                time.sleep(5)
    raise RuntimeError(f"Firebase {method} fail ho gaya: {last_err}")


def _discover_db_folders(home):
    """Folders jahan 'gfh database' (.xlsx/.xlsm/.csv) padi hai -
    Desktop / Downloads / Documents (2 level deep). User ka setup:
    credential.json usi folder me rakha hota hai jahan gfh database hai,
    is liye YE folder credential search me sab se pehle aate hain."""
    out = []
    seen = set()
    for sub in ("Desktop", "Downloads", "Documents"):
        base = os.path.join(home, sub)
        if not os.path.isdir(base):
            continue
        base_depth = base.rstrip(os.sep).count(os.sep)
        for root, dirs, files in os.walk(base):
            depth = root.rstrip(os.sep).count(os.sep) - base_depth
            if depth >= 2:
                dirs[:] = []          # isi se aage nahi
            if depth > 2:
                continue
            if root not in seen and find_default_file(root):
                seen.add(root)
                out.append(root)
                if len(out) >= 6:
                    return out
    return out


def find_credential(extra_folders=None):
    """Koi bhi SUPPORTED Firebase credential dhoondta hai (v3.7).
    Places: data-file folder (drag&drop / gfh database folder - AUTO),
    script folder, CWD, home, Desktop, Downloads, Documents + unke
    2-level subfolders. Names: *credential* (.json/.txt) + baqi .json.
    Prefer: service_account > authorized_user > web_config.
    Returns: (path, data, ctype, label, issue)
      - path/data/ctype/label : credential mili (best type)
      - issue                 : file mili thi par supported nahi (wajah)
      - sab None              : kahin nahi mili (kahan dekha - batata hai)"""
    def load_json(p):
        last_err = None
        for enc in ("utf-8-sig", "utf-16", "cp1252"):
            try:
                with open(p, "r", encoding=enc) as f:
                    return json.load(f), None
            except Exception as e:
                last_err = e
                continue
        return None, last_err

    def check_file(p):
        data, err = load_json(p)
        if data is None:
            return None, None, None, (f"{p} JSON read nahi hui "
                                      f"({str(err)[:80]})")
        d = unwrap_cred(data)
        ctype = classify_cred(d)
        if ctype:
            return data, ctype, CRED_LABELS[ctype], None
        keys = (", ".join(list(d.keys())[:8])
                if isinstance(d, dict) else type(d).__name__)
        return None, None, None, (f"{p} mili thi lekin ye koi supported "
                                  f"Firebase credential nahi "
                                  f"(andar ye keys thin: {keys})")

    home = os.path.expanduser("~")
    script_dir = os.path.dirname(os.path.abspath(__file__))
    # 1) data-file folder (jahan gfh database.xlsx hai) sab se pehle
    folders = [f for f in (extra_folders or []) if os.path.isdir(f)]
    # 2) khud dhoonde hue 'gfh database' wale folders
    folders += [f for f in _discover_db_folders(home) if f not in folders]
    # 3) standard jagahen
    folders += [f for f in ([script_dir, os.getcwd(), home]
                            + [os.path.join(home, sub) for sub in
                               ("Desktop", "Downloads", "Documents")])
                if f not in folders]
    home_areas = [os.path.join(home, sub)
                  for sub in ("Desktop", "Downloads", "Documents")]

    best = [None, None, None, None]      # path, data, ctype, label
    best_rank = 99
    first_issue = None

    def consider(p):
        nonlocal best_rank, first_issue
        data, ctype, label, issue = check_file(p)
        if ctype is not None:
            rank = {"service_account": 0, "authorized_user": 1,
                    "desktop_client": 2, "web_config": 3}[ctype]
            if rank < best_rank:
                best_rank = rank
                best[:] = [p, data, ctype, label]
            return rank == 0            # SA mil gaya -> scan khatam
        if issue and first_issue is None:
            first_issue = issue
        return False

    stop = False
    # PASS 0: pehle SAVE KIYA HUA login (credential_user.json) - ye
    # OAuth client file se ek dafa login ke baad banta hai; is se
    # dobara browser nahi khulta
    saved_login = os.path.join(script_dir, "credential_user.json")
    if os.path.isfile(saved_login):
        consider(saved_login)
    # PASS 1: *credential* naam wali files (sab folders, .txt variant samet)
    for folder in folders:
        if stop:
            break
        if not os.path.isdir(folder):
            continue
        try:
            names = sorted(os.listdir(folder))
        except Exception:
            continue
        for fn in names:
            low = fn.lower()
            if "credential" not in low or fn.startswith("~$"):
                continue
            if not (low.endswith(".json") or low.endswith(".txt")):
                continue
            if consider(os.path.join(folder, fn)):
                stop = True
                break

    # PASS 1b: SUBFOLDERS (2 level deep) - script folder, data-file
    # folder aur Desktop/Downloads/Documents sab ke andar bhi dekho
    walk_roots = [script_dir] + folders[:1] + home_areas
    seen_roots = set()
    for base in walk_roots:
        if stop:
            break
        if not base or not os.path.isdir(base) or base in seen_roots:
            continue
        seen_roots.add(base)
        base_depth = base.rstrip(os.sep).count(os.sep)
        for root, dirs, files in os.walk(base):
            if stop:
                break
            depth = root.rstrip(os.sep).count(os.sep) - base_depth
            if depth >= 2:
                dirs[:] = []      # isi se aage nahi
            if depth > 2:
                continue
            for fn in files:
                low = fn.lower()
                if "credential" not in low or fn.startswith("~$"):
                    continue
                if not (low.endswith(".json") or low.endswith(".txt")):
                    continue
                if consider(os.path.join(root, fn)):
                    stop = True
                    break

    # PASS 2: doosri .json files (alphabetical) - script + data folder
    if not stop:
        for folder in ([script_dir] + folders[:1]):
            try:
                names = sorted(os.listdir(folder))
            except Exception:
                continue
            for fn in names:
                if not fn.lower().endswith(".json") or fn.startswith("~$"):
                    continue
                if consider(os.path.join(folder, fn)):
                    stop = True
                    break
            if stop:
                break

    if best[0]:
        return best[0], best[1], best[2], best[3], None
    where = (", ".join(folders[:6]) or "kahin nahi")
    return None, None, None, None, (first_issue or
            f"in jagahon me supported credential nahi mili: {where}")


def authed_upload(node, payload, extra_folders=None):
    """v3.8 credential engine: type detect -> token mint -> authenticated
    DELETE + SINGLE PUT (dashboard wala shape) + shallow verify.
    desktop_client (gcloud OAuth client-secret) -> EK DAFA browser login.
    Returns: (True, None) kamyaab | (False, wajah) fail
             | (None, issue) supported credential nahi mili"""
    cred_path, cred, ctype, label, issue = find_credential(extra_folders)
    if not cred_path:
        return None, issue
    script_dir = os.path.dirname(os.path.abspath(__file__))
    if os.path.dirname(os.path.abspath(cred_path)) != script_dir:
        try:
            import shutil
            dst = os.path.join(script_dir, "credential.json")
            if not os.path.isfile(dst):
                shutil.copyfile(cred_path, dst)
                print(f"[OK] {os.path.basename(cred_path)} ka copy is folder "
                      f"mein credential.json naam se bana diya - agle run instant.")
        except Exception:
            pass

    print(f"[MODE] {label}")
    print(f"       File: {os.path.basename(cred_path)}")
    unwrapped = unwrap_cred(cred)
    pid = (unwrapped.get("project_id") or "").strip()
    if pid and pid != PROJECT_ID:
        print(f"[WARN] Credential ka project '{pid}' hai, expected "
              f"'{PROJECT_ID}' - phir bhi try karta hoon.")

    print("[AUTH] Google/Firebase token banata hoon...")
    if ctype == "desktop_client":
        try:
            token, _saved = desktop_login_flow(unwrapped, script_dir)
            auth = FBAuth("authorized_user", token)
        except Exception as e:
            msg = str(e)
            extra = ("\n        ALTERNATIVE (best, 2 min): Firebase Console > "
                     "gear icon > Project settings > Service accounts > "
                     "Generate new private key -> 'credential.json' IS "
                     "folder mein save karo.")
            return False, f"OAuth login fail: {msg[:250]}{extra}"
    else:
        try:
            auth = FBAuth(ctype, mint_token(cred, ctype))
        except Exception as e:
            msg = str(e)
            extra = ""
            if ctype == "web_config":
                extra = ("\n        FIX A: Firebase Console > Authentication > "
                         "Sign-in method > Anonymous > Enable -> dobara chalao.\n"
                         "        FIX B: Project settings > Service accounts > "
                         "Generate new private key -> 'credential.json' is folder mein.")
            return False, f"token nahi ban saka: {msg[:250]}{extra}"
    print("[OK] Authentication tayyar hai.")

    base = f"{FIREBASE_DB_URL}/{node}.json"
    print("[CLEAR] Purana data hata raha hoon (authenticated DELETE)...")
    try:
        status, resp = fb_request(auth.url(base), method="DELETE")
    except Exception as e:
        return False, f"clear (DELETE) fail: {str(e)[:200]}"
    if status in (401, 403):
        return False, (f"DELETE par HTTP {status} - is credential ko database "
                       f"par likhne ki ijazat nahi.\n        FIX: Firebase Console > "
                       f"Project settings > Service accounts > Generate new private "
                       f"key -> 'credential.json' is folder mein save karo "
                       f"(owner service-account rules bypass karti hai).")
    if status != 200:
        return False, f"clear (DELETE) fail: HTTP {status} {str(resp)[:200]}"
    print("[OK] Purana data clear ho gaya.")

    total = len(payload)
    print(f"[UPLOAD] {total:,} rows (header samet) ek hi PUT mein charhata hoon...")
    t0 = time.time()
    try:
        status, resp = fb_request(auth.url(base), method="PUT", payload=payload)
    except Exception as e:
        return False, f"upload (PUT) fail: {str(e)[:200]}"
    if status in (401, 403):
        # token expire ho sakta hai - ek dafa refresh kar ke dobara
        try:
            auth = FBAuth(ctype, mint_token(cred, ctype))
            status, resp = fb_request(auth.url(base), method="PUT", payload=payload)
        except Exception as e:
            return False, f"upload (PUT, refresh ke baad bhi) fail: {str(e)[:200]}"
    if status != 200:
        return False, f"upload (PUT) fail: HTTP {status} {str(resp)[:200]}"
    print(f"    [UPLOAD] {total:,}/{total:,} rows bhej diye "
          f"({time.time() - t0:.0f} sec)")

    try:
        vstatus, vdata = fb_request(
            auth.url(f"{FIREBASE_DB_URL}/{node}.json?shallow=true"))
        count = len(vdata) if (vstatus == 200 and isinstance(vdata, dict)) else -1
    except Exception:
        count = -1
    if count == total:
        print(f"[OK] Upload ho gaya + verify: {count:,} rows "
              f"({time.time() - t0:.0f} sec)")
        return True, None
    return False, f"verify fail: expected {total:,}, mila {count:,} - dobara chalao"


def upload(payload, dry_run=False, node=DATA_NODE, extra_folders=None):
    url = f"{FIREBASE_DB_URL}/{node}.json"
    size_mb = len(json.dumps(payload, ensure_ascii=False, separators=(",", ":"))) / 1024 / 1024
    n_rows = len(payload) - 1
    print(LINE)
    print(f"  File data   : {n_rows:,} data rows, {len(payload[0])} columns")
    print(f"  Upload size : {size_mb:.1f} MB")
    print(f"  Destination : {url}")
    print(LINE)

    if dry_run:
        print("[DRY-RUN] Upload NAHI hua - sirf preview. First data row:")
        print("  " + json.dumps(payload[1], ensure_ascii=False)[:500])
        return

    print("\n  Purana data delete ho kar is file ka data charhega - shuru...")

    ok, info = authed_upload(node, payload, extra_folders=extra_folders)
    if ok is True:
        print("\n  Dashboard kholo aur refresh karo:  https://gfhinventorydashboard.netlify.app")
        return
    if ok is None:
        print(LINE)
        if info:
            print(f"[ERROR] {info}")
        print("[ERROR] Supported credential nahi mili - upload ROAK diya.")
        print("        Rules LOCK hain, bina auth upload 401 deta hai -")
        print("        data corrupt nahi hoga, kuch bhi badla nahi.")
        print("        FIX (2 min): Firebase Console > gear icon > Project")
        print("        settings > Service accounts > Generate new private")
        print("        key > JSON download -> 'credential.json' naam se IS")
        print("        folder mein save karo -> RUN_UPLOAD.bat dobara.")
        print(LINE)
        return
    print(LINE)
    print(f"[ERROR] Upload fail ho gaya: {info}")
    print(LINE)
    return


def backup():
    base = f"{FIREBASE_DB_URL}/{DATA_NODE}.json"
    print("[BACKUP] Current dashboard data download kar raha hoon...")
    cred_path, cred, ctype, label, issue = find_credential()
    if not cred_path:
        print(f"[ERROR] Backup ke liye bhi credential chahiye.")
        print(f"        ({issue or 'supported credential nahi mili'})")
        return
    try:
        if ctype == "desktop_client":
            script_dir = os.path.dirname(os.path.abspath(__file__))
            token, _saved = desktop_login_flow(unwrap_cred(cred), script_dir)
            auth = FBAuth("authorized_user", token)
        else:
            auth = FBAuth(ctype, mint_token(cred, ctype))
    except Exception as e:
        print(f"[ERROR] Backup auth fail: {str(e)[:200]}")
        return
    status, data = fb_request(auth.url(base))
    if status != 200 or data is None:
        print(f"[ERROR] HTTP {status} - backup fail.")
        return
    rows = data if isinstance(data, list) else list(data.values())
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), f"GFH_backup_{ts}.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, separators=(",", ":"))
    print(f"[OK] Backup save hua: {os.path.basename(out)}  ({len(rows):,} rows)")
    print("     Ye file upload ke liye bhi use kar sakte ho (safety copy).")


def score_file_headers(path):
    """File ke headers dashboard ke 25 columns se kitne match karte hain -
    sirf pehli row peek (fast). xlsx mein sab sheets check hoti hain.
    Returns: score int (0-25), error par -1."""
    try:
        want_set = {h.lower() for h in EXPECTED_HEADERS}
        ext = os.path.splitext(path)[1].lower()
        if ext == ".csv":
            with open(path, "r", encoding="utf-8-sig", newline="") as f:
                header = next(csv.reader(f), [])
            got = [("" if h is None else str(h).strip().lower()) for h in header]
            return sum(1 for c in got if c in want_set)
        if not ensure_openpyxl():
            return -1
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        best = 0
        for ws in wb.worksheets:
            got = norm_first_row(ws)
            sc = sum(1 for c in got if c in want_set)
            best = max(best, sc)
        wb.close()
        return best
    except Exception:
        return -1


def find_best_file(folder):
    """Folder ki saari .xlsx/.xlsm/.csv ko headers match par score karta hai
    (SILENT - koi poori list print nahi hoti, v3.5 wala circus khatam).
    Best < 15 match ho to None (ghalat file auto-upload nahi hogi)."""
    cands = []
    for fn in os.listdir(folder):
        if fn.startswith("~$"):
            continue
        if fn.lower().endswith((".xlsx", ".xlsm", ".csv")):
            p = os.path.join(folder, fn)
            cands.append((os.path.getmtime(p), p))
    if not cands:
        return None
    scored = []
    for mt, p in sorted(cands, reverse=True):   # newest first
        scored.append((score_file_headers(p), mt, p))
    scored.sort(key=lambda x: (x[0], x[1]), reverse=True)  # score pehle, phir newest
    if scored[0][0] < AUTO_PICK_MIN:
        return None
    return scored[0][2]


LAST_FILE_MEM = "LAST_UPLOAD.txt"


def remember_last_file(path):
    """Pichli file ka naam memory file mein - agle run par zero scan."""
    try:
        with open(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                               LAST_FILE_MEM), "w", encoding="utf-8") as f:
            f.write(os.path.basename(path))
    except Exception:
        pass


def find_default_file(folder):
    """'gfh database.xlsx' naam ki file seedha uthata hai
    (case-insensitive, .xlsx/.xlsm/.csv - newest jeet ta hai)."""
    best = None
    try:
        for fn in os.listdir(folder):
            if fn.startswith("~$"):
                continue
            base, ext = os.path.splitext(fn.lower())
            if ext in (".xlsx", ".xlsm", ".csv") and base.strip() == "gfh database":
                p = os.path.join(folder, fn)
                if best is None or os.path.getmtime(p) > os.path.getmtime(best):
                    best = p
    except Exception:
        pass
    return best


def main():
    args = sys.argv[1:]
    print("=" * 58)
    print("   GFH INVENTORY DASHBOARD  -  FIREBASE UPLOADER v3.8")
    print("=" * 58)

    if "--backup" in args:
        backup()
        return

    node = DATA_NODE
    if "--node" in args:                       # hidden: testing only
        i = args.index("--node")
        node = args[i + 1]
        args = args[:i] + args[i + 2:]

    dry = "--dry-run" in args

    files = [a for a in args if not a.startswith("--")]

    if not files:
        folder = os.path.dirname(os.path.abspath(__file__))
        path = None
        # 1) memory: pichli baar jo file thi, wohi (ZERO scan)
        try:
            with open(os.path.join(folder, LAST_FILE_MEM), "r",
                      encoding="utf-8") as f:
                lastname = f.read().strip()
        except Exception:
            lastname = None
        if lastname:
            for fn in os.listdir(folder):
                if fn.lower() == lastname.lower() and not fn.startswith("~$"):
                    path = os.path.join(folder, fn)
                    print(f"[PICK] {fn}  (pichli baar wahi file upload hui thi)")
                    break
        # 2) default naam: gfh database.xlsx
        if path is None:
            path = find_default_file(folder)
            if path is not None:
                print(f"[PICK] {os.path.basename(path)}  (default 'gfh database' file)")
        # 3) ab bhi nahi -> chup-chaap scan, sirf winner print
        if path is None:
            path = find_best_file(folder)
            if path is not None:
                print(f"[PICK] {os.path.basename(path)}  (folder scan se best match)")
        if path is None:
            print("\n[ERROR] 'gfh database.xlsx' is folder mein nahi mili.")
            print("        File is folder mein rakho YA file ko RUN_UPLOAD.bat")
            print("        par DRAG & DROP karo - seedha upload ho jayega.")
            return
        files = [path]

    path = files[0]
    if not os.path.exists(path):
        print(f"[ERROR] File nahi mili: {path}")
        return
    if os.path.splitext(path)[1].lower() not in (".xlsx", ".xlsm", ".csv"):
        print(f"[ERROR] '{os.path.basename(path)}' upload ke liye sahi format nahi (.xlsx/.xlsm/.csv chahiye).")
        print("        Agar file .xls (purana format) hai to Excel mein Save As > .xlsx kar lo.")
        return

    remember_last_file(path)   # agli baar yehi file seedha pick hogi

    print(f"\n[STEP 1/2] File read + check kar raha hoon...")
    payload = build_payload(path)
    print("[STEP 2/2] Ready.")
    # credential bhi data-file ke folder me dekho (user ka setup: file
    # jahan gfh database hai wahi credential.json rakhi hoti hai)
    upload(payload, dry_run=dry, node=node,
           extra_folders=[os.path.dirname(os.path.abspath(path))])


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n[CANCEL] Band kar diya.")
